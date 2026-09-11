# -*- coding: utf-8 -*-
"""
问题 2 求解与结果输出：
  * 表 3 / 表 4（3 h 内每隔 0.5 h，距离 0/0.5/1/1.5/2 cm）
  * result2.xlsx（每隔 1 s、距离每隔 0.1 cm 的完整结果，两工作表：温度 / 水分浓度）
"""
import os
import sys
import time
import numpy as np

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
import p2_model as m

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
RES = os.path.join(BASE, "results")
os.makedirs(RES, exist_ok=True)

N = 20            # dr = 0.1 cm
DT = 1.0          # s
T_END = 259200.0  # 3 天（72 h，整个烘干过程）
R4 = "{:.4f}"


def main():
    T_inf, C_inf = m.make_ambient()
    dr = m.R_CYL / N

    t0 = time.time()
    it = []
    times, r, T, C = m.solve(T_inf, C_inf, t_end=T_END, dt=DT, N=N,
                             n_store=1, picard_max=3, newton=True, iters_out=it)
    print(f"[运行] 3 天模拟完成，耗时 {time.time()-t0:.1f} s，平均 Picard 迭代 {it[0]:.2f} 次")
    print(f"      帧数 {len(times)}，温度范围 [{T.min():.4f},{T.max():.4f}]，浓度范围 [{C.min():.4f},{C.max():.4f}]")

    r_cm = r * 100.0
    # ---- 表 3 / 表 4（3h 内每隔 0.5h，5 个距离） ----
    t_tab = np.array([1800, 3600, 5400, 7200, 9000, 10800])
    d_tab = [0.0, 0.5, 1.0, 1.5, 2.0]
    d_idx = [int(round(d / 0.1)) for d in d_tab]
    t_idx = [int(round(t / DT)) for t in t_tab]

    lines = []
    lines.append("表 3  3 小时内药材的温度（°C）")
    lines.append("时间/h\t" + "\t".join(f"{d:g} cm" for d in d_tab))
    for t, ti in zip(t_tab, t_idx):
        vals = [T[ti][di] for di in d_idx]
        lines.append(f"{t/3600:.1f}\t" + "\t".join(R4.format(v) for v in vals))
    lines.append("")
    lines.append("表 4  3 小时内药材的水分浓度（kg/kg，干基）")
    lines.append("时间/h\t" + "\t".join(f"{d:g} cm" for d in d_tab))
    for t, ti in zip(t_tab, t_idx):
        vals = [C[ti][di] for di in d_idx]
        lines.append(f"{t/3600:.1f}\t" + "\t".join(R4.format(v) for v in vals))
    table_text = "\n".join(lines)
    print("\n" + table_text)

    # ---- result2.xlsx ----
    _write_xlsx(times, r_cm, T, C)

    # ---- 关键结论 ----
    idx = np.where(C[:, 0] < 0.15)[0]
    t_dry = times[idx[0]] if len(idx) else None
    lines.append("")
    if t_dry is not None:
        lines.append(f"中心含水率降至 0.15 kg/kg 的时刻 t = {t_dry:.0f} s = {t_dry/3600:.1f} h（烘干时长）")
    lines.append(f"3 天终态：中心 T={T[-1][0]:.4f} °C、C={C[-1][0]:.4f} kg/kg；表面 C={C[-1][-1]:.4f} kg/kg")

    out_txt = os.path.join(RES, "problem2_tables.txt")
    with open(out_txt, "w", encoding="utf-8") as f:
        f.write("\n".join(lines) + "\n")
    print(f"\n已写入 {out_txt}")


def _write_xlsx(times, r_cm, T, C):
    """写出 result2.xlsx（时间 1s 起，距离 0--2 cm 每 0.1 cm）。"""
    import openpyxl
    from openpyxl.utils import get_column_letter

    out = os.path.join(RES, "result2.xlsx")
    wb = openpyxl.Workbook(write_only=True)

    dist = [f"{v:.1f}" for v in r_cm]  # 0.0, 0.1, ..., 2.0

    for sheet_name, field in (("温度", T), ("水分浓度", C)):
        ws = wb.create_sheet(title=sheet_name)
        ws.append(["时间\\到药材中心的距离"] + dist)
        # 时间从 1 s 起，每隔 1 s
        for k in range(1, len(times)):
            row = [times[k]] + [round(float(v), 4) for v in field[k]]
            ws.append(row)

    wb.save(out)
    print(f"已写入 {out}（{len(times)-1} 行 × {len(dist)} 列 × 2 工作表）")


if __name__ == "__main__":
    main()
