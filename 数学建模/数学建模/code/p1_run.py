# -*- coding: utf-8 -*-
"""
问题 1 主驱动程序：
    * 读取附件 1（烘房温度 / 水分浓度随时间变化）
    * 调用 Crank-Nicolson 求解器求解一维非稳态热质耦合模型
    * 输出题目要求的表 1、表 2
    * 生成 result1.xlsx（温度、水分浓度两个工作表，1800 s × 21 个半径位置）
"""
import os
import sys
import numpy as np
import openpyxl
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from p1_model import solve, R_CYL, T0, C0

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
ATT = os.path.join(os.path.dirname(BASE), "A题", "附件")
RES = os.path.join(BASE, "results")
TEMPLATE = os.path.join(ATT, "附件3", "result1.xlsx")

DT = 1.0          # 时间步长 s
N = 20            # 径向网格数 -> dr = 2cm/20 = 0.1 cm
T_END = 1800.0
PICARD_MAX = 3   # D(C) 滞后迭代次数（≥2 即可恢复时间二阶精度）


# ----------------------------------------------------------------------------
# 附件 1：烘房环境条件
# ----------------------------------------------------------------------------
def load_ambient(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["Sheet1"]
    rows = [r for r in ws.iter_rows(values_only=True)][1:]
    t = np.array([float(r[0]) for r in rows])
    T = np.array([float(r[1]) for r in rows])
    C = np.array([float(r[2]) for r in rows])
    return t, T, C


def make_linear(t_data, y_data):
    """构造分段线性插值函数（区间外按端点常数延拓）。"""
    def f(t):
        return float(np.interp(t, t_data, y_data))
    return f


def main():
    t_amb, T_amb, C_amb = load_ambient(os.path.join(ATT, "附件1.xlsx"))
    T_inf_fn = make_linear(t_amb, T_amb)
    C_inf_fn = make_linear(t_amb, C_amb)

    print(f"环境数据: t ∈ [{t_amb[0]:.0f}, {t_amb[-1]:.0f}] s, {len(t_amb)} 点")
    print(f"烘房温度   T_inf: {T_amb[0]:.3f} -> {T_amb[-1]:.3f} degC")
    print(f"烘房水分浓度 C_inf: {C_amb[0]:.5f} -> {C_amb[-1]:.5f} kg/kg")
    print(f"网格: N={N}, dr={R_CYL/N*100:.2f} cm, dt={DT} s, CN(theta=0.5)\n")

    times, r, T_hist, C_hist = solve(
        T_inf_fn, C_inf_fn, t_end=T_END, dt=DT, N=N, theta=0.5, n_store=1,
        bc="halfcell", picard_max=PICARD_MAX
    )

    r_cm = r * 100.0                       # 到药材中心的距离 (cm)
    dist_out = np.round(np.arange(0, 21) * 0.1, 1)     # 0.0 ... 2.0 cm

    # ------------------------------------------------------------------
    # 表 1 / 表 2
    # ------------------------------------------------------------------
    t_tab = [100, 300, 600, 900, 1200, 1500, 1800]
    d_tab = [0.0, 0.5, 1.0, 1.5, 2.0]

    lines = []
    lines.append("表 1  30 分钟内药材的温度 (degC)")
    lines.append("时间/s\t" + "\t".join(f"{d:g} cm" for d in d_tab))
    for ts in t_tab:
        i = int(ts / DT)
        vals = [T_hist[i][int(round(d / (R_CYL / N * 100)))] for d in d_tab]
        lines.append(f"{ts}\t" + "\t".join(f"{v:.4f}" for v in vals))

    lines.append("")
    lines.append("表 2  30 分钟内药材的水分浓度 (kg/kg)")
    lines.append("时间/s\t" + "\t".join(f"{d:g} cm" for d in d_tab))
    for ts in t_tab:
        i = int(ts / DT)
        vals = [C_hist[i][int(round(d / (R_CYL / N * 100)))] for d in d_tab]
        lines.append(f"{ts}\t" + "\t".join(f"{v:.4f}" for v in vals))

    table_text = "\n".join(lines)
    print(table_text)
    with open(os.path.join(RES, "problem1_tables.txt"), "w", encoding="utf-8") as f:
        f.write(table_text + "\n")

    # ------------------------------------------------------------------
    # result1.xlsx —— 基于模板填充
    # ------------------------------------------------------------------
    wb = openpyxl.load_workbook(TEMPLATE)
    for sheet, hist, name in (("温度", T_hist, "温度"), ("水分浓度", C_hist, "水分浓度")):
        ws = wb[sheet]
        # 第 1 行表头: B..V 列对应 0, 0.1, ..., 2.0 cm
        for k, d in enumerate(dist_out):
            ws.cell(row=1, column=2 + k, value=float(d))
        for i in range(1, int(T_END / DT) + 1):
            ws.cell(row=1 + i, column=1, value=i)
            for k, d in enumerate(dist_out):
                j = int(round(d / (R_CYL / N * 100)))    # 网格索引
                ws.cell(row=1 + i, column=2 + k, value=round(float(hist[i][j]), 4))
    out_xlsx = os.path.join(RES, "result1.xlsx")
    wb.save(out_xlsx)
    print(f"\n已写出 {out_xlsx}  (时间 1..1800 s, 距离 0..2.0 cm)")

    np.savez(os.path.join(RES, "p1_fields.npz"),
             times=times, r=r, T=T_hist, C=C_hist)
    print("已写出 results/p1_fields.npz")


if __name__ == "__main__":
    main()
