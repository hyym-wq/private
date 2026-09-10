from pathlib import Path
import importlib.util,json
import numpy as np
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[2]
spec=importlib.util.spec_from_file_location('solver',Path(__file__).with_name('问题一_数值比较.py'))
m=importlib.util.module_from_spec(spec);spec.loader.exec_module(m)
r,dx,faces,w=m.geometry(6)
op=m.operator(np.linspace(1e-9,7e-9,7),8e-7,dx,faces,w)
A=np.diag(op[1])+np.diag(op[0][1:],-1)+np.diag(op[2][:-1],1)
b=np.linspace(1,2,7)
err=float(np.max(np.abs(m.linear_solve(m.factor(op,.25,.5),b)-np.linalg.solve(np.eye(7)-.125*A,b))))
eq=float(np.max(np.abs(m.apply(op,np.full(7,2.55),2.55))))
assert err<1e-12 and eq<1e-12
p=ROOT/'04_结果/问题一';report=json.loads((p/'统计与验证.json').read_text(encoding='utf-8'))
for c in report['runs'].values():
 assert c['finite'] and c['bounded'] and c['radial']
 assert c['min_temperature_increment']>=-1e-9 and c['max_moisture_increment']<=1e-9
 if c['method']=='CN':assert c['max_residual']<1e-12
 else:assert c['max_explicit_loss']<=1
arr=np.load(p/'数值场.npz');book=p/'result1.xlsx'
if book.exists():
 wb=load_workbook(book,data_only=False)
 assert wb.sheetnames==['温度','水分浓度']
 for name,key in [('温度','T'),('水分浓度','C')]:
  s=wb[name];assert (s.max_row,s.max_column)==(1801,22)
  a=np.array(list(s.values)[1:],float)
  assert np.array_equal(a[:,0],np.arange(1,1801))
  assert np.array_equal(a[:,1:],np.round(arr[key][1:,::4],4))
  assert all(s.cell(i,j).number_format=='0.0000' for i in range(2,1802) for j in range(2,23))
 wb.close()
else:print('尚未导出Excel，仅验证数值结果')
m.dump(p/'独立检查.json',{'三对角求解与numpy稠密求解最大差':err,'均匀平衡场右端残差':eq,'结果检查通过':True,'Excel已检查':book.exists()})
print('独立求解器、均匀场、趋势、残差与现有Excel检查通过')
