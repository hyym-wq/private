from pathlib import Path
import argparse, csv, hashlib, json, time
import numpy as np
from openpyxl import load_workbook

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'04_结果/问题一'
FIG=ROOT/'05_图片/问题一'
R=0.02
RHO,CP,K,H,HM=820.,2600.,0.36,25.,8e-7
ALPHA=K/(RHO*CP)
TIMES=np.arange(1801)
SELECT=[100,300,600,900,1200,1500,1800]

def dump(p,obj):
    p.parent.mkdir(parents=True,exist_ok=True)
    p.write_text(json.dumps(obj,ensure_ascii=False,indent=2,allow_nan=False),encoding='utf-8')

def csvout(p,header,rows):
    with p.open('w',encoding='utf-8-sig',newline='') as f:
        w=csv.writer(f);w.writerow(header);w.writerows(rows)

def audit(path,numeric=True):
    w=load_workbook(path,data_only=False)
    info={'文件':path.name,'SHA256':hashlib.sha256(path.read_bytes()).hexdigest(),'表':[]}
    arrays=[]
    for s in w:
        rows=list(s.values)
        record={'名称':s.title,'状态':s.sheet_state,'尺寸':[s.max_row,s.max_column],
                '表头':list(rows[0]),'合并':str(s.merged_cells),'非空数':sum(v is not None for row in rows for v in row)}
        if numeric:
            if any(v is None or isinstance(v,bool) or not isinstance(v,(float,int)) for row in rows[1:] for v in row):
                raise ValueError('输入缺失或非数值:'+path.name)
            a=np.array(rows[1:],float)
            assert np.isfinite(a).all() and np.all(np.diff(a[:,0])>0)
            record.update({'缺失':0,'列类型':['数值']*a.shape[1], 'min':a.min(0).tolist(),'max':a.max(0).tolist()})
            arrays.append(a)
        else:
            record['单元格']=rows
        info['表'].append(record)
    w.close()
    return info,arrays

def dcoef(c):
    if np.any(c<=0): raise ValueError('非正浓度，不进行裁剪')
    return 7e-9*np.exp(-0.89/c)

def geometry(N):
    r=np.linspace(0,R,N+1); dx=R/N
    faces=np.r_[0,(r[:-1]+r[1:])/2,R]
    w=np.diff(faces**2)/2
    return r,dx,faces,w

def operator(a,b,dx,faces,w):
    g=faces[1:-1]*(a[1:]+a[:-1])/(2*dx)
    lo=np.r_[0,g/w[1:]];up=np.r_[g/w[:-1],0]
    source=np.zeros(len(a));source[-1]=R*b/w[-1]
    return lo,-lo-up-source,up,source

def apply(op,u,ambient):
    lo,di,up,src=op
    v=di*u+src*ambient
    v[1:]+=lo[1:]*u[:-1];v[:-1]+=up[:-1]*u[1:]
    return v

def factor(op,h,theta):
    lo,di,up,_=op
    lower=-theta*h*lo;diag=1-theta*h*di;upper=-theta*h*up
    diag=diag.copy();mult=np.zeros(len(diag))
    for i in range(1,len(diag)):
        mult[i]=lower[i]/diag[i-1];diag[i]-=mult[i]*upper[i-1]
    return mult,diag,upper

def linear_solve(fac,rhs):
    mult,diag,upper=fac
    x=rhs.copy()
    for i in range(1,len(x)):x[i]-=mult[i]*x[i-1]
    x[-1]/=diag[-1]
    for i in range(len(x)-2,-1,-1):x[i]=(x[i]-upper[i]*x[i+1])/diag[i]
    return x

def solve(N,dt,method,boundary):
    start=time.perf_counter()
    r,dx,faces,w=geometry(N);weights=w/w.sum()
    nt=round(1800/dt);stride=round(1/dt)
    assert abs(nt*dt-1800)<1e-9 and abs(stride*dt-1)<1e-9
    ts=np.arange(nt+1)*dt
    env=np.column_stack([np.interp(ts,boundary[:,0],boundary[:,j]) for j in [1,2]])
    T=np.full(N+1,28.);C=np.full(N+1,2.55)
    Tout=np.empty((1801,N+1));Cout=Tout.copy();Tout[0]=T;Cout[0]=C
    hop=operator(np.full(N+1,ALPHA),H/(RHO*CP),dx,faces,w)
    heat_factor=factor(hop,dt,.5)
    max_iter=0;max_res=0.;max_loss=0.;balanceT=0.;balanceC=0.
    min_dT=0.;max_dC=0.;finite=True;bounded=True;radial=True
    iter_total=0
    for n in range(nt):
        # 第一完整时间步采用两个后向Euler半步；其余采用真正CN。
        steps=[(dt/2,1.,ts[n],ts[n]+dt/2),(dt/2,1.,ts[n]+dt/2,ts[n+1])] if method=='CN' and n==0 else [(dt,0. if method=='Euler' else .5,ts[n],ts[n+1])]
        for h,theta,t0,t1 in steps:
            oldT=T.copy();oldC=C.copy()
            e0=np.array([np.interp(t0,boundary[:,0],boundary[:,j]) for j in [1,2]])
            e1=np.array([np.interp(t1,boundary[:,0],boundary[:,j]) for j in [1,2]])
            cop=operator(dcoef(C),HM,dx,faces,w)
            fT=apply(hop,T,e0[0]);fC=apply(cop,C,e0[1])
            if method=='Euler':
                loss=h*max((-hop[1]).max(),(-cop[1]).max());max_loss=max(max_loss,float(loss))
                if loss>1+1e-12:raise ValueError('Euler不稳定')
                T=T+h*fT;C=C+h*fC
                it=0
            else:
                fac=heat_factor if theta==.5 and h==dt else factor(hop,h,theta)
                T=linear_solve(fac,oldT+h*(1-theta)*fT+h*theta*hop[3]*e1[0])
                guess=C.copy()
                for it in range(1,31):
                    newop=operator(dcoef(guess),HM,dx,faces,w)
                    rhs=oldC+h*(1-theta)*fC+h*theta*newop[3]*e1[1]
                    new=linear_solve(factor(newop,h,theta),rhs)
                    trueop=operator(dcoef(new),HM,dx,faces,w)
                    residual=new-oldC-h*((1-theta)*fC+theta*apply(trueop,new,e1[1]))
                    res=float(np.max(np.abs(residual)))
                    if np.max(np.abs(new-guess))<1e-11 and res<1e-12:
                        C=new;max_res=max(max_res,res);break
                    guess=new
                else:raise RuntimeError('Picard迭代未收敛')
            max_iter=max(max_iter,it);iter_total+=it
            qT=R*H/(RHO*CP)*((1-theta)*(e0[0]-oldT[-1])+theta*(e1[0]-T[-1]))
            qC=R*HM*((1-theta)*(e0[1]-oldC[-1])+theta*(e1[1]-C[-1]))
            balanceT=max(balanceT,abs(float(np.dot(w,T-oldT)-h*qT))*2*np.pi*RHO*CP)
            balanceC=max(balanceC,abs(float(np.dot(w,C-oldC)-h*qC))/w.sum())
            min_dT=min(min_dT,float(np.min(T-oldT)));max_dC=max(max_dC,float(np.max(C-oldC)))
            finite &= bool(np.isfinite(T).all() and np.isfinite(C).all())
            bounded &= bool(T.min()>=28-1e-9 and T.max()<=e1[0]+1e-9 and C.min()>=e1[1]-1e-9 and C.max()<=2.55+1e-9)
            radial &= bool(np.diff(T).min()>=-1e-9 and np.diff(C).max()<=1e-9)
        if (n+1)%stride==0:
            idx=(n+1)//stride;Tout[idx]=T;Cout[idx]=C
    check={'method':method,'N':N,'nodes':N+1,'dt':dt,'seconds':time.perf_counter()-start,
           'max_iterations':max_iter,'iterations_total':iter_total,'max_residual':max_res,
           'max_explicit_loss':max_loss,'heat_balance_J_m':balanceT,'moisture_balance':balanceC,
           'min_temperature_increment':min_dT,'max_moisture_increment':max_dC,
           'finite':bool(finite),'bounded':bool(bounded),'radial':bool(radial),
           'mean_C_final':float(np.dot(weights,C)),'T_final':T[::max(1,N//4)].tolist(),'C_final':C[::max(1,N//4)].tolist()}
    if not (finite and bounded and radial and min_dT>=-1e-9 and max_dC<=1e-9):raise RuntimeError(str(check))
    for key,arr in [('T',Tout),('C',Cout)]:
        check[key+'_stats']={'min':float(arr.min()),'max':float(arr.max()),'mean':float(arr.mean()),'std':float(arr.std()),'CV':float(arr.std()/arr.mean()),'amplitude':float(np.ptp(arr))}
    return dict(T=Tout,C=Cout,r=r,check=check)

def delta(a,b):
    step=(len(b['r'])-1)//(len(a['r'])-1)
    return {k:float(np.max(np.abs(a[k]-b[k][:,::step]))) for k in ['T','C']}

def generate_figures(runs,boundary,report):
    import matplotlib
    matplotlib.use('Agg')
    import matplotlib.pyplot as plt
    plt.rcParams.update({'font.sans-serif':['Microsoft YaHei'],'axes.unicode_minus':False,'font.size':10,'figure.dpi':160,'savefig.dpi':160})
    main=runs['CN_80_0.25']; radius=main['r']*100
    def save(fig,name):
        fig.tight_layout();fig.savefig(FIG/(name+'.png'),bbox_inches='tight');plt.close(fig)
    for key,label,cmap in [('T','温度（℃）','inferno'),('C','干基含水率（kg/kg）','viridis')]:
        fig,ax=plt.subplots(figsize=(6.5,3.6))
        im=ax.pcolormesh(radius,TIMES,main[key],shading='auto',cmap=cmap,rasterized=True)
        ax.set(xlabel='到中心轴的距离（cm）',ylabel='时间（s）');fig.colorbar(im,ax=ax,label=label)
        save(fig,'温度时空分布' if key=='T' else '水分时空分布')
    fig,axes=plt.subplots(1,2,figsize=(10,3.5))
    for t in [100,600,1200,1800]:
        for ax,key in zip(axes,['T','C']):ax.plot(radius,main[key][t],label=f'{t} s',lw=1.6)
    for ax,label in zip(axes,['温度（℃）','干基含水率（kg/kg）']):
        ax.set(xlabel='到中心轴的距离（cm）',ylabel=label);ax.legend(fontsize=8);ax.grid(alpha=.25)
    save(fig,'径向剖面')
    fig,axes=plt.subplots(1,2,figsize=(10,3.5))
    for i,label in [(0,'中心'),(40,'半径1 cm'),(80,'表面')]:
        for ax,key in zip(axes,['T','C']):ax.plot(TIMES,main[key][:,i],label=label,lw=1.6)
    axes[0].plot(boundary[:31,0],boundary[:31,1],'--',label='烘房',lw=1.6)
    for ax,label in zip(axes,['温度（℃）','干基含水率（kg/kg）']):ax.set(xlabel='时间（s）',ylabel=label);ax.legend(fontsize=8);ax.grid(alpha=.25)
    save(fig,'时间变化')
    fig,axes=plt.subplots(1,2,figsize=(10,3.5))
    eu=runs['Euler_40_0.25'];cn=runs['CN_40_0.25']
    for ax,key,label in zip(axes,['T','C'],['温度最大绝对差（℃）','含水率最大绝对差（kg/kg）']):
        ax.plot(TIMES,np.max(np.abs(eu[key]-cn[key]),axis=1),lw=1.6)
        ax.set(xlabel='时间（s）',ylabel=label);ax.grid(alpha=.25)
    save(fig,'Euler与CN差异')
    fig,axes=plt.subplots(2,2,figsize=(9,6))
    for col,(key,label) in enumerate([('T','温度最大绝对差（℃）'),('C','含水率最大绝对差（kg/kg）')]):
        axes[0,col].plot([1,.5],[report['time_1_05'][key],report['time_05_025'][key]],marker='o',lw=1.6)
        axes[1,col].plot([1,.5],[report['space_20_40'][key],report['space_40_80'][key]],marker='o',lw=1.6)
        for row,xlabel in enumerate(['较粗时间步（s）','较粗网格间距（mm）']):
            axes[row,col].set(xlabel=xlabel,ylabel=label,yscale='log');axes[row,col].grid(alpha=.25)
    save(fig,'收敛检查')

def main():
    parser=argparse.ArgumentParser();parser.add_argument('--no-figures',action='store_true');args=parser.parse_args()
    for p in [OUT,FIG,ROOT/'work',ROOT/'02_数据/问题一']:p.mkdir(parents=True,exist_ok=True)
    inputs=ROOT/'02_数据/问题一'
    if not (inputs/'附件1.xlsx').exists():
        import shutil
        source=ROOT/'00_赛题与附件/附件'
        for name in ['附件1.xlsx','附件2.xlsx']:shutil.copy2(source/name,inputs/name)
        shutil.copy2(source/'附件3/result1.xlsx',inputs/'result1模板.xlsx')
    a1,arrays=audit(inputs/'附件1.xlsx');a2,_=audit(inputs/'附件2.xlsx');a3,_=audit(inputs/'result1模板.xlsx',False)
    boundary=arrays[0]
    assert boundary.shape==(241,3) and boundary[0,0]==0 and boundary[-1,0]>=1800
    dump(OUT/'输入核查.json',[a1,a2,a3])
    csvout(inputs/'边界预处理.csv',['时间_s','烘房温度_C','烘房水分浓度_kg_kg'],zip(TIMES,np.interp(TIMES,boundary[:,0],boundary[:,1]),np.interp(TIMES,boundary[:,0],boundary[:,2])))
    configs=[('Euler',40,.25),('CN',40,1.),('CN',40,.5),('CN',40,.25),('CN',20,.25),('CN',80,.25)]
    runs={}
    for method,N,dt in configs:
        key=f'{method}_{N}_{dt:g}';run=solve(N,dt,method,boundary);runs[key]=run
        print(key,json.dumps(run['check'],ensure_ascii=False),flush=True)
    report={
      'equal_grid_methods':delta(runs['Euler_40_0.25'],runs['CN_40_0.25']),
      'time_1_05':delta(runs['CN_40_1'],runs['CN_40_0.5']),
      'time_05_025':delta(runs['CN_40_0.5'],runs['CN_40_0.25']),
      'space_20_40':delta(runs['CN_20_0.25'],runs['CN_40_0.25']),
      'space_40_80':delta(runs['CN_40_0.25'],runs['CN_80_0.25'])}
    report['runs']={key:r['check'] for key,r in runs.items()}
    mainrun=runs['CN_80_0.25']
    # 同时间步的细网格结果是数值参考，不是解析解。
    report['Euler_to_fine']=delta(runs['Euler_40_0.25'],mainrun)
    report['CN_to_fine']=delta(runs['CN_40_0.25'],mainrun)
    dump(OUT/'统计与验证.json',report)
    for key in ['T','C']:
        a=mainrun[key]
        csvout(OUT/('指定时刻温度.csv' if key=='T' else '指定时刻水分浓度.csv'),['时间_s','0_cm','0.5_cm','1_cm','1.5_cm','2_cm'],[[t]+[f'{v:.4f}' for v in a[t,::20]] for t in SELECT])
    np.savez_compressed(OUT/'数值场.npz',time=TIMES,r=mainrun['r'],T=mainrun['T'],C=mainrun['C'],Euler_T=runs['Euler_40_0.25']['T'],Euler_C=runs['Euler_40_0.25']['C'],comparison_CN_T=runs['CN_40_0.25']['T'],comparison_CN_C=runs['CN_40_0.25']['C'])
    dump(ROOT/'work/表格数据.json',{'模板':str(inputs/'result1模板.xlsx'),'半径_cm':np.round(mainrun['r'][::4]*100,1).tolist(),
         '温度':np.column_stack([TIMES[1:],np.round(mainrun['T'][1:,::4],4)]).tolist(),
         '水分浓度':np.column_stack([TIMES[1:],np.round(mainrun['C'][1:,::4],4)]).tolist()})
    if not args.no_figures:generate_figures(runs,boundary,report)
    print('COMPARISON',json.dumps({k:v for k,v in report.items() if k!='runs'},ensure_ascii=False),flush=True)

if __name__=='__main__':main()
