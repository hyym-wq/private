"""2026-09-10 第一问结果表只读核验 v01 Codex.

Compare every saved XLSX numerical result with the canonical solver JSON.
openpyxl is used strictly for reading; workbook authoring is artifact-tool JS.
"""
from __future__ import annotations

import hashlib
import json
import math
from pathlib import Path

import openpyxl


def main() -> None:
    private_dir = Path(__file__).resolve().parents[2]
    result_dir = private_dir / "04_结果" / "第一问"
    source_path = result_dir / "q1_export.json"
    workbook_path = result_dir / "result1.xlsx"
    source = json.loads(source_path.read_text(encoding="utf-8"))
    workbook = openpyxl.load_workbook(workbook_path, data_only=False)
    expected_sheets = [("温度", "temperature"), ("水分浓度", "moisture")]
    assert workbook.sheetnames == [x[0] for x in expected_sheets], workbook.sheetnames
    report = {"status": "pass", "source": source_path.name,
              "source_run": source.get("source_run"),
              "cross_section": source.get("cross_section"),
              "workbook": workbook_path.name, "comparison_tolerance": 1e-12,
              "sha256": {p.name: hashlib.sha256(p.read_bytes()).hexdigest()
                         for p in (source_path, workbook_path)}, "sheets": {}}
    for sheet_name, key in expected_sheets:
        sheet = workbook[sheet_name]
        assert (sheet.max_row, sheet.max_column) == (1801, 22)
        assert sheet.freeze_panes == "B2", sheet.freeze_panes
        assert len(sheet.merged_cells.ranges) == 0
        assert "(s)" in sheet["A1"].value and "(cm)" in sheet["A1"].value
        headers = [sheet.cell(1, i + 2).value for i in range(21)]
        assert all(isinstance(x, (int, float)) for x in headers)
        assert all(abs(a - b) < 1e-12 for a, b in zip(headers, source["radii_cm"]))
        maximum_difference = 0.0
        numeric_results = 0
        for i in range(1800):
            assert sheet.cell(i + 2, 1).value == source["times"][i] == i + 1
            for j in range(21):
                cell = sheet.cell(i + 2, j + 2)
                value = cell.value
                assert isinstance(value, (int, float)) and not isinstance(value, bool)
                assert math.isfinite(value)
                assert cell.data_type == "n"
                assert cell.number_format == "0.0000", cell.number_format
                delta = abs(value - source[key][i][j])
                assert delta <= 1e-12, (sheet_name, cell.coordinate, delta)
                maximum_difference = max(maximum_difference, delta)
                numeric_results += 1
        report["sheets"][sheet_name] = {
            "dimensions": [1801, 22], "time_range_s": [1, 1800],
            "radius_range_cm": [0, 2], "radius_step_cm": 0.1,
            "numeric_results": numeric_results, "max_abs_difference": maximum_difference,
            "data_number_format": "0.0000", "freeze_panes": sheet.freeze_panes,
        }
    report["notes"] = [
        "Canonical JSON and XLSX numerical values are rounded to four decimal places as required; unrounded solver arrays are retained in the NPZ results.",
        "No Excel formulas or scientific model recalculation is needed for this numerical-results export.",
        "Mathematical accuracy and convergence are established by the solver review, separately from export checks.",
    ]
    output_path = result_dir / "0910_第一问_结果表核验_v01_Codex.json"
    output_path.write_text(json.dumps(report, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(report, ensure_ascii=True, indent=2))


if __name__ == "__main__":
    main()
