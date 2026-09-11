"""Q3: continue the approved Q2 model and detect max(C) < 0.15.

Usage: python q3_compute.py --N 20 --dt 1
The N=20, dt=1 run is the Q2-compatible result; other grids are diagnostics.
"""
from pathlib import Path
import argparse
import hashlib
import importlib.util
import json
import platform
import time

import numpy as np
import openpyxl

ROOT = Path(__file__).resolve().parents[2]
MODEL = ROOT / "数学建模/数学建模/code/p2_model.py"
ATT = ROOT.parent / "CUMCM2026Problems/A题/附件"
OUT = ROOT / "results/Q3"


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("--N", type=int, default=20)
    parser.add_argument("--dt", type=float, default=1)
    parser.add_argument("--export", action="store_true", help="Export this refined grid at the required 0.1 cm spacing")
    args = parser.parse_args()
    spec = importlib.util.spec_from_file_location("q2_inherited", MODEL)
    model = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(model)
    ambient = ATT / "附件1.xlsx"
    ta, te, ca = model.load_ambient(ambient)
    assert np.isfinite(np.stack([ta, te, ca])).all()
    assert np.all(np.diff(ta) > 0) and ta[0] == 0
    assert args.N % 20 == 0 and 60 / args.dt == int(60 / args.dt)
    Tinf, Cinf = model.make_ambient(ambient)
    start = time.perf_counter()
    iterations = []
    times, radius, T, C = model.solve(
        Tinf, Cinf, t_end=259200, dt=args.dt, N=args.N,
        n_store=1, picard_max=3, newton=True, iters_out=iterations)
    elapsed = time.perf_counter() - start
    maxima = C.max(axis=1)
    eligible = np.flatnonzero(maxima < 0.15)
    assert len(eligible), "No drying event within 72 hours"
    end = int(eligible[0])
    previous = end - 1
    cross = times[previous] + args.dt * (maxima[previous] - 0.15) / (maxima[previous] - maxima[end])
    table_times = list(np.arange(21600., times[end], 21600.)) + [times[end]]
    table_idx = np.rint(np.array(table_times) / args.dt).astype(int)
    radial_idx = np.array([0, 1, 2, 3, 4]) * (args.N // 4)
    stride = int(round(60 / args.dt))
    sample_idx = list(range(stride, end + 1, stride))
    if sample_idx[-1] != end:
        sample_idx.append(end)
    sample_idx = np.asarray(sample_idx)
    out_radius_idx = np.arange(21) * (args.N // 20)
    OUT.mkdir(parents=True, exist_ok=True)
    tag = f"N{args.N}_dt{args.dt:g}"
    np.savez_compressed(OUT / f"fields_{tag}.npz", times=times[sample_idx],
                        radius=radius, C=C[sample_idx], T=T[sample_idx],
                        table_times=np.array(table_times), table_C=C[table_idx][:, radial_idx])
    summary = {
        "schema_version": 1, "question": "Q3", "approved_decision_id": "q3_inherit_q2",
        "status": "computed", "N": args.N, "dt_s": args.dt,
        "role": "refined_main" if args.export else ("q2_compatible" if args.N == 20 and args.dt == 1 else "grid_diagnostic"),
        "runtime_s": elapsed, "mean_iterations": iterations[0],
        "drying_time_s": float(times[end]), "drying_time_h": float(times[end] / 3600),
        "threshold_crossing_interpolated_s": float(cross),
        "previous_time_s": float(times[previous]), "previous_max_C": float(maxima[previous]),
        "terminal_max_C": float(maxima[end]), "terminal_surface_C": float(C[end, -1]),
        "maximum_off_center_excess": float(np.max(maxima[:end+1] - C[:end+1, 0])),
        "maximum_radial_increase": float(np.max(np.diff(C[:end+1], axis=1))),
        "C_min_before_stop": float(C[:end+1].min()),
        "C_max_before_stop": float(C[:end+1].max()),
        "finite_fields": bool(np.isfinite(T).all() and np.isfinite(C).all()),
        "table_times_h": (np.asarray(table_times) / 3600).tolist(),
        "table_C": C[table_idx][:, radial_idx].tolist(),
        "sample_rows": len(sample_idx), "terminal_row_added": bool(times[end] % 60 != 0),
        "ambient_last_time_s": float(ta[-1]), "ambient_last_T_C": float(te[-1]),
        "ambient_last_C": float(ca[-1]),
        "inputs": {str(p): hashlib.sha256(p.read_bytes()).hexdigest() for p in [MODEL, ambient]},
        "environment": {"python": platform.python_version(), "numpy": np.__version__, "openpyxl": openpyxl.__version__},
        "assumptions": ["Inherited Q2 1D constant-radius model", "Ambient held at attachment endpoint after 4 h", "No stage reset at 1800 s", "Inherited max 3 nonlinear iterations; iteration convergence not certified per step"],
        "seed": None,
    }
    if (args.N == 20 and args.dt == 1) or args.export:
        wb = openpyxl.load_workbook(ATT / "附件3/result3.xlsx")
        ws = wb.active
        ws.delete_rows(1, ws.max_row)
        ws.append(["时间\\到药材中心的距离"] + [round(float(x * 100), 1) for x in radius[out_radius_idx]])
        for i in sample_idx:
            ws.append([float(times[i])] + [round(float(x), 4) for x in C[i, out_radius_idx]])
        wb.save(OUT / "result3.xlsx")
        lines = ["时间/h,0 cm,0.5 cm,1 cm,1.5 cm,2 cm"]
        for t, row in zip(table_times, C[table_idx][:, radial_idx]):
            lines.append(f"{t/3600:.4f}," + ",".join(f"{x:.4f}" for x in row))
        (OUT / "table5.csv").write_text("\n".join(lines) + "\n", encoding="utf-8-sig")
        check = openpyxl.load_workbook(OUT / "result3.xlsx", read_only=True, data_only=True)
        rows = list(check.active.values)
        assert len(rows) == len(sample_idx) + 1 and len(rows[0]) == 22
        assert rows[-1][0] == times[end]
        for row, i in zip(rows[1:], sample_idx):
            assert row[0] == times[i]
            assert np.allclose(row[1:], np.round(C[i, out_radius_idx], 4), atol=1e-12, rtol=0)
        check.close()
        summary["xlsx_readback"] = "PASS"
    (OUT / f"run_summary_{tag}.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8")
    print(json.dumps(summary, ensure_ascii=False, indent=2), flush=True)


if __name__ == "__main__":
    main()
