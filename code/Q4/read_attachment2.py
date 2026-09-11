from pathlib import Path
import zipfile
import xml.etree.ElementTree as ET
import openpyxl
root=Path(__file__).resolve().parents[2]
files=list((root.parent/'CUMCM2026Problems/A题/附件').glob('*'))
print(files)
for p in files:
 print(p.name,p.is_file())
 if '附件2' in p.name:
  with zipfile.ZipFile(p) as z:
   print(z.namelist())
  wb=openpyxl.load_workbook(p,data_only=True)
  print(wb.sheetnames)
  for ws in wb.worksheets:
   print(ws.title,ws.max_row,ws.max_column,list(ws.values)[:8])
if True:
 p=root.parent/'CUMCM2026Problems/A题/附件/附件3/result4.xlsx'
 wb=openpyxl.load_workbook(p,data_only=True)
 print('result4',wb.sheetnames)
 for ws in wb.worksheets: print(ws.title,ws.max_row,ws.max_column,list(ws.values)[:8])
